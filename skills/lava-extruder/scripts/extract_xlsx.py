#!/usr/bin/env python3
"""
extract_xlsx.py — Extract Excel spreadsheets (.xlsx, .xls) to structured Markdown.
Usage: extract_xlsx.py <input_file> [output.md]

Features:
1. Renders each worksheet as a separate '## Sheet Name' section.
2. Formats tabular data into clean GitHub-Flavored Markdown tables.
3. Supports openpyxl, pandas, or standard-library XML zip parsing for .xlsx.
4. Falls back to LibreOffice headless conversion if python libraries are unavailable.
"""

import sys
import os
import zipfile
import xml.etree.ElementTree as ET
import subprocess
import tempfile


def format_markdown_table(rows: list[list[str]]) -> str:
    """Format a 2D list into an aligned Markdown table."""
    # Filter empty rows
    clean_rows = []
    for r in rows:
        if any(c.strip() for c in r):
            clean_rows.append([c.strip().replace("\n", " ") for c in r])

    if not clean_rows:
        return ""

    col_count = max(len(r) for r in clean_rows)
    # Ensure uniform columns
    padded_rows = [r + [""] * (col_count - len(r)) for r in clean_rows]

    col_widths = [max(len(str(r[i])) for r in padded_rows) for i in range(col_count)]
    col_widths = [max(w, 3) for w in col_widths]  # Minimum width 3 for ---

    lines = []
    header = padded_rows[0]
    lines.append("| " + " | ".join(header[i].ljust(col_widths[i]) for i in range(col_count)) + " |")
    lines.append("| " + " | ".join("-" * col_widths[i] for i in range(col_count)) + " |")

    for r in padded_rows[1:]:
        lines.append("| " + " | ".join(r[i].ljust(col_widths[i]) for i in range(col_count)) + " |")

    return "\n".join(lines) + "\n\n"


def convert_via_openpyxl(input_path: str, output_path: str) -> bool:
    try:
        import openpyxl
        wb = openpyxl.load_workbook(input_path, data_only=True, read_only=True)
        sections = [f"# {os.path.splitext(os.path.basename(input_path))[0]}\n"]

        for sheet_name in wb.sheetnames:
            sheet = wb[sheet_name]
            rows = []
            for row in sheet.iter_rows(values_only=True):
                str_row = ["" if val is None else str(val) for val in row]
                rows.append(str_row)
            table_md = format_markdown_table(rows)
            if table_md.strip():
                sections.append(f"## {sheet_name}\n\n" + table_md)

        wb.close()
        if len(sections) > 1:
            with open(output_path, "w", encoding="utf-8") as f:
                f.write("\n\n".join(sections).strip() + "\n")
            return True
    except Exception:
        pass
    return False


def convert_via_pandas(input_path: str, output_path: str) -> bool:
    try:
        import pandas as pd
        sheets = pd.read_excel(input_path, sheet_name=None)
        sections = [f"# {os.path.splitext(os.path.basename(input_path))[0]}\n"]

        for sheet_name, df in sheets.items():
            if df.empty:
                continue
            rows = [list(df.columns)]
            for _, r in df.iterrows():
                rows.append([str(v) if pd.notna(v) else "" for v in r])
            table_md = format_markdown_table(rows)
            if table_md.strip():
                sections.append(f"## {sheet_name}\n\n" + table_md)

        if len(sections) > 1:
            with open(output_path, "w", encoding="utf-8") as f:
                f.write("\n\n".join(sections).strip() + "\n")
            return True
    except Exception:
        pass
    return False


def convert_via_stdlib_xlsx(input_path: str, output_path: str) -> bool:
    """Zero-dependency parser for .xlsx (openxml zip + sharedStrings.xml)."""
    if not input_path.lower().endswith(".xlsx"):
        return False

    try:
        with zipfile.ZipFile(input_path, "r") as z:
            # Read shared strings
            shared_strings = []
            if "xl/sharedStrings.xml" in z.namelist():
                ss_tree = ET.fromstring(z.read("xl/sharedStrings.xml"))
                ns = {"main": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
                for si in ss_tree.findall(".//main:si", ns) or ss_tree.findall(".//si"):
                    t_nodes = si.findall(".//main:t", ns) or si.findall(".//t")
                    shared_strings.append("".join(t.text or "" for t in t_nodes))

            # Read workbook.xml to get sheet names
            wb_tree = ET.fromstring(z.read("xl/workbook.xml"))
            sheets = []
            for s in wb_tree.findall(".//{http://schemas.openxmlformats.org/spreadsheetml/2006/main}sheet"):
                sheets.append((s.attrib.get("name", "Sheet"), s.attrib.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id", "")))

            # Read worksheets
            sections = [f"# {os.path.splitext(os.path.basename(input_path))[0]}\n"]
            sheet_files = sorted([n for n in z.namelist() if n.startswith("xl/worksheets/sheet") and n.endswith(".xml")])

            for idx, sfile in enumerate(sheet_files):
                sheet_name = sheets[idx][0] if idx < len(sheets) else f"Sheet {idx + 1}"
                s_tree = ET.fromstring(z.read(sfile))
                rows = []
                for row_el in s_tree.findall(".//{http://schemas.openxmlformats.org/spreadsheetml/2006/main}row"):
                    row_cells = []
                    for c in row_el.findall(".//{http://schemas.openxmlformats.org/spreadsheetml/2006/main}c"):
                        c_type = c.attrib.get("t")
                        v = c.find("{http://schemas.openxmlformats.org/spreadsheetml/2006/main}v")
                        val = v.text if v is not None and v.text else ""
                        if c_type == "s" and val.isdigit():
                            idx_ss = int(val)
                            val = shared_strings[idx_ss] if idx_ss < len(shared_strings) else val
                        row_cells.append(val)
                    if row_cells:
                        rows.append(row_cells)

                table_md = format_markdown_table(rows)
                if table_md.strip():
                    sections.append(f"## {sheet_name}\n\n" + table_md)

            if len(sections) > 1:
                with open(output_path, "w", encoding="utf-8") as f:
                    f.write("\n\n".join(sections).strip() + "\n")
                return True
    except Exception:
        pass
    return False


def convert_via_libreoffice(input_path: str, output_path: str) -> bool:
    if subprocess.run(["which", "libreoffice"], capture_output=True).returncode != 0:
        return False

    with tempfile.TemporaryDirectory() as tmpdir:
        cmd = ["libreoffice", "--headless", "--convert-to", "csv", "--outdir", tmpdir, input_path]
        subprocess.run(cmd, capture_output=True)
        csv_name = os.path.splitext(os.path.basename(input_path))[0] + ".csv"
        csv_path = os.path.join(tmpdir, csv_name)
        if os.path.exists(csv_path):
            # Read CSV and format as markdown table
            import csv
            rows = []
            with open(csv_path, "r", encoding="utf-8", errors="ignore") as f:
                reader = csv.reader(f)
                for r in reader:
                    rows.append(r)
            table_md = format_markdown_table(rows)
            if table_md.strip():
                with open(output_path, "w", encoding="utf-8") as f:
                    f.write(f"# {os.path.splitext(os.path.basename(input_path))[0]}\n\n" + table_md)
                return True
    return False


def main():
    if len(sys.argv) < 2:
        print("Usage: extract_xlsx.py <input_file> [output.md]", file=sys.stderr)
        sys.exit(1)

    input_path = sys.argv[1]
    output_path = sys.argv[2] if len(sys.argv) > 2 else f"{input_path}.extracted.md"

    if not os.path.isfile(input_path):
        print(f"Error: File not found: {input_path}", file=sys.stderr)
        sys.exit(1)

    # Try methods in order
    if convert_via_openpyxl(input_path, output_path):
        lines = len(open(output_path).readlines())
        print(f"Done. Extracted {lines} lines to {output_path} via openpyxl", file=sys.stderr)
        return

    if convert_via_pandas(input_path, output_path):
        lines = len(open(output_path).readlines())
        print(f"Done. Extracted {lines} lines to {output_path} via pandas", file=sys.stderr)
        return

    if convert_via_stdlib_xlsx(input_path, output_path):
        lines = len(open(output_path).readlines())
        print(f"Done. Extracted {lines} lines to {output_path} via stdlib XML", file=sys.stderr)
        return

    if convert_via_libreoffice(input_path, output_path):
        lines = len(open(output_path).readlines())
        print(f"Done. Extracted {lines} lines to {output_path} via LibreOffice", file=sys.stderr)
        return

    print("Error: Could not extract spreadsheet.", file=sys.stderr)
    print("Install openpyxl for best results: pip install openpyxl", file=sys.stderr)
    sys.exit(1)


if __name__ == "__main__":
    main()
